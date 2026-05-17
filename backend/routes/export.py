"""数据导出路由：CSV / JSON / Excel（records/income/assets）+ 月报 HTML。"""
from __future__ import annotations

import csv
import io
import json

from flask import Blueprint, Response, abort, g, jsonify, request, stream_with_context

from auth import login_required
from db import get_db

export_bp = Blueprint("export", __name__)

# UTF-8 BOM 让 Excel 正确识别中文
_UTF8_BOM = "\ufeff"

_KIND_QUERIES = {
    "records": (
        "SELECT date, category, amount, note FROM records WHERE user_id = ? "
        "ORDER BY date DESC, id DESC",
        ["日期", "分类", "金额", "备注"],
    ),
    "income": (
        "SELECT date, category, amount, note FROM income WHERE user_id = ? "
        "ORDER BY date DESC, id DESC",
        ["日期", "来源", "金额", "备注"],
    ),
    "assets": (
        "SELECT name, type, symbol, holdings, cost_basis, current_value, currency, notes "
        "FROM assets WHERE user_id = ? ORDER BY current_value DESC",
        ["名称", "类型", "代码", "数量", "成本", "现值", "币种", "备注"],
    ),
}


@export_bp.route("/api/export/csv", methods=["GET"])
@login_required
def export_csv():
    """兼容旧 URL：保留 /api/export/csv 端点，等价于 /api/export?format=csv。"""
    return _do_export("csv")


@export_bp.route("/api/export", methods=["GET"])
@login_required
def export_any():
    """统一导出入口：?type=records|income|assets & ?format=csv|json|xlsx。"""
    return _do_export((request.args.get("format") or "csv").strip().lower())


def _do_export(fmt: str) -> Response:
    kind = (request.args.get("type") or "records").strip()
    if kind not in _KIND_QUERIES:
        abort(400, description=f"不支持的导出类型：{kind}")
    if fmt not in ("csv", "json", "xlsx"):
        abort(400, description=f"不支持的导出格式：{fmt}")

    sql, headers = _KIND_QUERIES[kind]
    user_id = g.user_id

    if fmt == "csv":
        return _stream_csv(kind, sql, headers, user_id)
    if fmt == "json":
        return _emit_json(kind, sql, headers, user_id)
    return _emit_xlsx(kind, sql, headers, user_id)


# ---------------- CSV：流式输出，避免大账户内存峰值 ----------------

def _stream_csv(kind, sql, headers, user_id) -> Response:
    def _generate():
        buf = io.StringIO()
        writer = csv.writer(buf)

        def _flush():
            data = buf.getvalue()
            buf.seek(0)
            buf.truncate(0)
            return data

        buf.write(_UTF8_BOM)
        writer.writerow(headers)
        yield _flush()

        cursor = get_db().execute(sql, (user_id,))
        try:
            for r in cursor:
                writer.writerow([r[k] if r[k] is not None else "" for k in r.keys()])
                yield _flush()
        finally:
            cursor.close()

    return Response(
        stream_with_context(_generate()),
        mimetype="text/csv; charset=utf-8",
        headers={
            "Content-Disposition": f'attachment; filename="{kind}.csv"',
        },
    )


# ---------------- JSON：一次性返回结构化数组 ----------------

def _emit_json(kind, sql, headers, user_id) -> Response:
    rows = [dict(r) for r in get_db().execute(sql, (user_id,)).fetchall()]
    payload = {"type": kind, "count": len(rows), "fields": headers, "rows": rows}
    body = json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")
    return Response(
        body,
        mimetype="application/json; charset=utf-8",
        headers={
            "Content-Disposition": f'attachment; filename="{kind}.json"',
        },
    )


# ---------------- Excel：openpyxl 写一份带表头样式的 .xlsx ----------------

def _emit_xlsx(kind, sql, headers, user_id) -> Response:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        abort(501, description="缺少 openpyxl 依赖，无法导出 Excel；请联系管理员")

    wb = Workbook()
    ws = wb.active
    ws.title = kind

    # 表头样式：暖米底、棕褐字、加粗 + 居中
    header_fill = PatternFill(fgColor="F0ECE2", fill_type="solid")
    header_font = Font(name="Calibri", bold=True, color="794F27")
    center = Alignment(horizontal="center", vertical="center")
    for col, label in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center

    cursor = get_db().execute(sql, (user_id,))
    try:
        row_idx = 2
        for r in cursor:
            for col, key in enumerate(r.keys(), start=1):
                v = r[key]
                ws.cell(row=row_idx, column=col, value=v)
            row_idx += 1
    finally:
        cursor.close()

    # 自动列宽（按列内容长度估算）
    for col, label in enumerate(headers, start=1):
        max_len = len(str(label))
        for r in range(2, ws.max_row + 1):
            v = ws.cell(row=r, column=col).value
            if v is not None:
                max_len = max(max_len, min(40, len(str(v))))
        ws.column_dimensions[chr(64 + col)].width = max(10, min(40, max_len + 2))

    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        buf.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{kind}.xlsx"',
        },
    )


@export_bp.route("/api/export/report.html", methods=["GET"])
@login_required
def export_report_html():
    """把指定月份的报告渲染成可打印 HTML（用户在浏览器中 Cmd+P 转 PDF）。"""
    period = (request.args.get("period") or "").strip()
    if len(period) != 7:
        abort(400, description="period 参数应为 YYYY-MM")

    from services.reports import get_report
    rep = get_report(g.user_id, period)
    if not rep:
        abort(404, description="该月报告尚未生成")

    content = rep.get("content") or ""
    html = _render_report_html(period, rep.get("created_at") or "", content)
    return Response(html, mimetype="text/html; charset=utf-8")


def _render_report_html(period: str, created_at: str, markdown: str) -> str:
    """极简 Markdown→HTML 渲染（标题/列表/段落/强调），用于打印。"""
    import html as _html

    lines = markdown.split("\n")
    out: list[str] = []
    in_list = False

    def close_list():
        nonlocal in_list
        if in_list:
            out.append("</ul>")
            in_list = False

    for line in lines:
        s = line.rstrip()
        if s.startswith("### "):
            close_list(); out.append(f"<h3>{_html.escape(s[4:])}</h3>"); continue
        if s.startswith("## "):
            close_list(); out.append(f"<h2>{_html.escape(s[3:])}</h2>"); continue
        if s.startswith("# "):
            close_list(); out.append(f"<h1>{_html.escape(s[2:])}</h1>"); continue
        if s.startswith("- ") or s.startswith("* "):
            if not in_list:
                out.append("<ul>"); in_list = True
            item = _html.escape(s[2:])
            item = _bold(item)
            out.append(f"<li>{item}</li>"); continue
        close_list()
        if not s.strip():
            out.append(""); continue
        out.append(f"<p>{_bold(_html.escape(s))}</p>")
    close_list()
    body = "\n".join(out)

    return f"""<!doctype html>
<html lang="zh"><head><meta charset="utf-8">
<title>{period} 月度报告</title>
<style>
  body {{ font-family: -apple-system, "PingFang SC", "Microsoft YaHei", sans-serif;
         max-width: 760px; margin: 32px auto; padding: 0 24px; color: #1f2937; line-height: 1.7; }}
  h1 {{ font-size: 24px; border-bottom: 2px solid #3B82F6; padding-bottom: 6px; }}
  h2 {{ font-size: 18px; color: #2563EB; margin-top: 22px; }}
  h3 {{ font-size: 15px; margin-top: 16px; }}
  ul {{ padding-left: 22px; }}
  .meta {{ color: #6b7280; font-size: 12px; margin-bottom: 18px; }}
  @media print {{
    .no-print {{ display: none; }}
    body {{ margin: 0; padding: 0 12mm; }}
  }}
  .no-print {{ position: fixed; right: 16px; top: 16px;
              background: #3B82F6; color: #fff; border: none;
              padding: 8px 14px; border-radius: 6px; cursor: pointer; }}
</style></head>
<body>
  <button class="no-print" onclick="window.print()">🖨️ 打印 / 另存为 PDF</button>
  <h1>📑 {period} 月度报告</h1>
  <div class="meta">生成时间：{_html.escape(created_at)}</div>
  {body}
</body></html>"""


def _bold(s: str) -> str:
    import re as _re
    s = _re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", s)
    s = _re.sub(r"`([^`]+)`", r"<code>\1</code>", s)
    return s
